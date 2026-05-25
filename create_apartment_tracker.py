import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side,
    GradientFill
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles.numbers import FORMAT_DATE_DATETIME
import datetime

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Apartment Tracker"

# ── Colour palette ──────────────────────────────────────────────────────────
NAVY        = "1F3864"   # header row background
GOLD        = "C9A84C"   # header text
LIGHT_BLUE  = "DCE6F1"   # even row shading
WHITE       = "FFFFFF"
GREEN_FILL  = "E2EFDA"   # purchase rows
YELLOW_FILL = "FFF2CC"   # rental rows
ORANGE_FILL = "FCE4D6"   # passed rows

# ── Column definitions ────────────────────────────────────────────────────
columns = [
    ("Address",            28),
    ("Unit",               10),
    ("StreetEasy Link",    35),
    ("Type",               12),   # Rent / Purchase
    ("Asking Price",       15),
    ("Offer Price",        15),
    ("Broker Name",        22),
    ("Broker Email",       30),
    ("Date of Showing",    18),
    ("Status",             18),
    ("Notes",              50),
]

# ── Header row ───────────────────────────────────────────────────────────
header_font    = Font(name="Calibri", bold=True, color=GOLD, size=11)
header_fill    = PatternFill("solid", fgColor=NAVY)
header_align   = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin           = Side(style="thin", color="AAAAAA")
thin_border    = Border(left=thin, right=thin, top=thin, bottom=thin)

ws.row_dimensions[1].height = 32
for col_idx, (col_name, col_width) in enumerate(columns, start=1):
    cell = ws.cell(row=1, column=col_idx, value=col_name)
    cell.font    = header_font
    cell.fill    = header_fill
    cell.alignment = header_align
    cell.border  = thin_border
    ws.column_dimensions[get_column_letter(col_idx)].width = col_width

# ── Freeze the header row + add auto-filter ──────────────────────────────
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"

# ── Apartment data ───────────────────────────────────────────────────────
# (address, unit, streeteasy_link, type, asking, offer, broker, email, showing_date, status, notes)
apartments = [
    (
        "200 East 90th Street, New York, NY 10128",
        "16E",
        "https://streeteasy.com/building/200-east-90th-street-new_york",
        "Purchase",
        "$1,150,000",
        "$1,020,000",
        "Sean Devine (Devine Team at Compass)",
        "sean.devine@compass.com",
        "May 3, 2026",
        "Active – Negotiating",
        (
            "Co-op. Asking reduced from $1.25M → $1.15M. "
            "Offer $1,020,000; seller countered at $1,125,000. "
            "High floor, East River views. "
            "Buyer's broker: Saree Ptak (sareerptak@gmail.com)."
        ),
    ),
    (
        "10 East End Avenue, New York, NY 10028",
        "18A",
        "https://streeteasy.com/building/10-east-end-avenue-new_york",
        "Purchase",
        "$1,150,000",
        "$1,150,000",
        "Jonathan Tauzowicz (The Collective Team at Compass)",
        "jt@compass.com",
        "Jan 11, 2026",
        "Under Contract",
        (
            "Co-op. Estate sale (Estate of Tuppatsch). "
            "Agreed price $1,150,000. Contract in progress with attorney Andy Cutler. "
            "Court action (estate) adjourned to 5/5/26. "
            "Chase mortgage (7yr ARM 4.750%). "
            "Buyer's broker: Saree Ptak (sareerptak@gmail.com)."
        ),
    ),
    (
        "520 East 72nd Street, New York, NY 10021",
        "PHC",
        "https://streeteasy.com/building/520-east-72nd-street-new_york",
        "Purchase",
        "",
        "",
        "Louise Devlin (Brown Harris Stevens)",
        "LDevlin@bhsusa.com",
        "Jan 11, 2026",
        "Inquired",
        (
            "Co-op. 30% down building; no W/D in unit; no investors allowed. "
            "DTI < 30% required. 2% buyer's agent commission. "
            "Buyer's broker: Saree Ptak (sareerptak@gmail.com)."
        ),
    ),
    (
        "165 East 72nd Street, New York, NY 10021",
        "12J",
        "https://streeteasy.com/building/165-east-72nd-street-new_york",
        "Purchase",
        "",
        "",
        "Elisabeth Quick (Serhant)",
        "quick@serhant.com",
        "Dec 2025",
        "Passed",
        (
            "Passed – renovation cost too high. "
            "Buyer's broker: Saree Ptak (sareerptak@gmail.com)."
        ),
    ),
    (
        "310 East 70th Street, New York, NY 10021",
        "10E",
        "https://streeteasy.com/building/310-east-70th-street-new_york",
        "Purchase",
        "$1,395,000",
        "$1,225,000",
        "Jesse Klein (Compass)",
        "jesse.klein@compass.com",
        "Nov 5, 2025",
        "Passed",
        (
            "Co-op. Asking reduced from ~$1.45M → $1,395,000. "
            "Offer $1,225,000; seller countered at $1,395,000. "
            "Subletting policy: 2 out of 5 years. "
            "Buyer's broker: Saree Ptak (sareerptak@gmail.com)."
        ),
    ),
    (
        "239 East 79th Street, New York, NY 10075",
        "13L",
        "https://streeteasy.com/building/239-east-79th-street-new_york",
        "Purchase",
        "$1,375,000",
        "",
        "Roger J. Gillen (Brown Harris Stevens)",
        "RGillen@bhsusa.com",
        "Nov 2025",
        "Passed",
        (
            "Co-op. Asking $1,375,000. Passed – major renovation needed (2011 updates outdated). "
            "Broker rejected low offer. "
            "Buyer's broker: Saree Ptak (sareerptak@gmail.com)."
        ),
    ),
    (
        "420 East 80th Street, New York, NY 10075",
        "PHD",
        "https://streeteasy.com/for-rent/listing/420-east-80th-street",
        "Rent",
        "",
        "",
        "Maria Skovranova (The Ashford NYC)",
        "Leasing_Office-theashfordnyc@knck.io",
        "May 2025",
        "Inquired",
        (
            "Rental inquiry via StreetEasy May 2025. "
            "Building: The Ashford NYC. Contact Maria Skovranova. "
            "Price not confirmed."
        ),
    ),
]

# ── Row styles keyed by type / status ────────────────────────────────────
def row_fill(apt_type, status):
    if "passed" in status.lower():
        return PatternFill("solid", fgColor=ORANGE_FILL)
    if apt_type.lower() == "purchase":
        return PatternFill("solid", fgColor=GREEN_FILL)
    return PatternFill("solid", fgColor=YELLOW_FILL)

data_font      = Font(name="Calibri", size=10)
data_align_ww  = Alignment(vertical="top", wrap_text=True)
data_align_c   = Alignment(horizontal="center", vertical="top")
link_font      = Font(name="Calibri", size=10, color="0563C1", underline="single")

for row_idx, apt in enumerate(apartments, start=2):
    (address, unit, se_link, apt_type, asking, offer,
     broker, email, showing, status, notes) = apt

    fill = row_fill(apt_type, status)
    ws.row_dimensions[row_idx].height = 60

    values = [address, unit, se_link, apt_type, asking, offer,
              broker, email, showing, status, notes]

    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.font   = data_font
        cell.fill   = fill
        cell.border = thin_border

        if col_idx == 3 and value:          # StreetEasy link – make it a hyperlink
            cell.font      = link_font
            cell.hyperlink = value
            cell.value     = value
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        elif col_idx in (1, 7, 11):         # wrap text cols
            cell.alignment = data_align_ww
        elif col_idx in (2, 4, 5, 6, 9, 10):
            cell.alignment = data_align_c
        else:
            cell.alignment = data_align_ww

# ── Instructions / Legend tab ─────────────────────────────────────────────
ws2 = wb.create_sheet("How to Use")
ws2.column_dimensions["A"].width = 80

instructions = [
    ("NYC APARTMENT TRACKER – HOW TO USE", True),
    ("", False),
    ("EDITING FROM EXCEL", True),
    ("• Open NYC_Apartment_Tracker.xlsx directly and edit any cell.", False),
    ("• To add a new apartment: copy any existing row and paste at the bottom, then fill in the fields.", False),
    ("• Save when done. The file is stored in the repo.", False),
    ("", False),
    ("EDITING FROM CLAUDE", True),
    ('  Example: "Add 123 Main St, Unit 4A, asking $1.2M, purchase, broker Jane Smith at jane@compass.com, showing June 1"', False),
    ('  Example: "Update the notes for 10 East End Ave 18A"', False),
    ("• Claude reads this file, applies your change, saves, and commits.", False),
    ("", False),
    ("UPLOADING A SCREENSHOT OR EMAIL TO CHECK IF IT'S IN THE LIST", True),
    ("• Paste or drag an image / screenshot into the Claude chat.", False),
    ('  Ask: "Is this apartment already in my tracker?"', False),
    ("• Claude will read the address from the image and search the spreadsheet for a match.", False),
    ("", False),
    ("STATUS KEY", True),
    ("• Active – Negotiating  →  offer submitted, back-and-forth in progress", False),
    ("• Under Contract        →  agreed price, attorney review / board package stage", False),
    ("• Inquired              →  reached out, waiting or toured only", False),
    ("• Passed                →  decided not to proceed", False),
    ("", False),
    ("ROW COLOUR KEY", True),
    ("• Green  →  Purchase listing, active or under contract", False),
    ("• Yellow →  Rental listing", False),
    ("• Orange →  Passed / no longer pursuing", False),
]

title_font = Font(name="Calibri", bold=True, size=12)
body_font  = Font(name="Calibri", size=10)

for r_idx, (text, is_title) in enumerate(instructions, start=1):
    cell = ws2.cell(row=r_idx, column=1, value=text)
    cell.font = title_font if is_title else body_font
    cell.alignment = Alignment(wrap_text=True)

output_path = "/home/user/Default/NYC_Apartment_Tracker.xlsx"
wb.save(output_path)
print(f"Saved → {output_path}")
print(f"Apartments logged: {len(apartments)}")
